"""
FUTURE_DS_02: Customer Retention & Churn Analysis
Module: Excel Executive Workbook Builder
Author: Future Interns - Data Science & Analytics Track
"""

import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def build_retention_excel_workbook(clean_path: str, output_path: str):
    print(f"[*] Reading cleaned dataset from: {clean_path}")
    df = pd.read_csv(clean_path)

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    light_blue_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # 1. Executive KPIs Sheet
    ws1 = wb.create_sheet(title="Executive_KPIs")
    ws1.views.sheetView[0].showGridLines = True
    ws1["A1"] = "CUSTOMER RETENTION & CHURN AUDIT - EXECUTIVE SCORECARD"
    ws1["A1"].font = title_font

    kpis = [
        ("Total Subscribed Customer Accounts", len(df), "#,##0"),
        ("Retained Active Customers", (df['Churn'] == 'No').sum(), "#,##0"),
        ("Churned Customers (Lost Accounts)", (df['Churn'] == 'Yes').sum(), "#,##0"),
        ("Overall Customer Churn Rate", (df['Churn'] == 'Yes').mean(), "0.0%"),
        ("Active Customer Retention Rate", (df['Churn'] == 'No').mean(), "0.0%"),
        ("Average Account Tenure (Months)", df['Tenure_Months'].mean(), "0.0"),
        ("Average Monthly Charges ($)", df['MonthlyCharges'].mean(), "$#,##0.00"),
        ("Total Monthly Recurring Revenue (MRR)", df['MonthlyCharges'].sum(), "$#,##0.00"),
        ("Monthly Revenue Lost to Churn (MRR Risk)", df[df['Churn'] == 'Yes']['MonthlyCharges'].sum(), "$#,##0.00"),
        ("Average Customer Lifetime Value (CLV)", df['Customer_Lifetime_Value'].mean(), "$#,##0.00")
    ]

    ws1.cell(row=3, column=1, value="Core Retention Performance Metric").fill = navy_fill
    ws1.cell(row=3, column=1).font = header_font
    ws1.cell(row=3, column=2, value="Realized Value").fill = navy_fill
    ws1.cell(row=3, column=2).font = header_font

    for i, (kpi, val, num_fmt) in enumerate(kpis, start=4):
        c1 = ws1.cell(row=i, column=1, value=kpi)
        c2 = ws1.cell(row=i, column=2, value=val)
        c1.border = thin_border
        c2.border = thin_border
        c2.number_format = num_fmt
        c2.alignment = Alignment(horizontal="right")
        if i % 2 == 1:
            c1.fill = light_blue_fill

    # 2. Contract Analysis Sheet
    ws2 = wb.create_sheet(title="Contract_Analysis")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "CHURN RATE AUDIT BY CONTRACT TYPE"
    ws2["A1"].font = title_font

    ct_df = df.groupby('Contract_Type').agg(
        Total=('CustomerID', 'count'),
        Churned=('Churn_Binary', 'sum'),
        Avg_Monthly=('MonthlyCharges', 'mean'),
        Avg_Tenure=('Tenure_Months', 'mean')
    ).reset_index()
    ct_df['Churn_Rate'] = ct_df['Churned'] / ct_df['Total']

    headers2 = ["Contract Type", "Total Accounts", "Churned Accounts", "Churn Rate", "Avg Monthly Spend", "Avg Tenure (Mos)"]
    for col_idx, h in enumerate(headers2, 1):
        c = ws2.cell(row=3, column=col_idx, value=h)
        c.fill = navy_fill
        c.font = header_font

    for r_idx, row in enumerate(ct_df.itertuples(), start=4):
        ws2.cell(row=r_idx, column=1, value=row.Contract_Type).border = thin_border
        ws2.cell(row=r_idx, column=2, value=row.Total).border = thin_border
        ws2.cell(row=r_idx, column=3, value=row.Churned).border = thin_border
        c4 = ws2.cell(row=r_idx, column=4, value=row.Churn_Rate)
        c4.border = thin_border
        c4.number_format = "0.0%"
        c5 = ws2.cell(row=r_idx, column=5, value=row.Avg_Monthly)
        c5.border = thin_border
        c5.number_format = "$#,##0.00"
        c6 = ws2.cell(row=r_idx, column=6, value=round(row.Avg_Tenure, 1))
        c6.border = thin_border
        c6.number_format = "0.0"

    # 3. Tenure Brackets Sheet
    ws3 = wb.create_sheet(title="Tenure_Cohorts")
    ws3.views.sheetView[0].showGridLines = True
    ws3["A1"] = "RETENTION PERFORMANCE BY TENURE COHORT"
    ws3["A1"].font = title_font

    tb_df = df.groupby('Tenure_Bracket', observed=False).agg(
        Total=('CustomerID', 'count'),
        Churned=('Churn_Binary', 'sum'),
        Total_Revenue=('TotalCharges', 'sum')
    ).reset_index()
    tb_df['Churn_Rate'] = tb_df['Churned'] / tb_df['Total']

    headers3 = ["Tenure Bracket", "Total Accounts", "Churned Accounts", "Cohort Churn Rate", "Realized Spend ($)"]
    for col_idx, h in enumerate(headers3, 1):
        c = ws3.cell(row=3, column=col_idx, value=h)
        c.fill = navy_fill
        c.font = header_font

    for r_idx, row in enumerate(tb_df.itertuples(), start=4):
        ws3.cell(row=r_idx, column=1, value=str(row.Tenure_Bracket)).border = thin_border
        ws3.cell(row=r_idx, column=2, value=row.Total).border = thin_border
        ws3.cell(row=r_idx, column=3, value=row.Churned).border = thin_border
        c4 = ws3.cell(row=r_idx, column=4, value=row.Churn_Rate)
        c4.border = thin_border
        c4.number_format = "0.0%"
        c5 = ws3.cell(row=r_idx, column=5, value=row.Total_Revenue)
        c5.border = thin_border
        c5.number_format = "$#,##0.00"

    # 4. Cleaned Data Sample (Top 500 rows)
    ws4 = wb.create_sheet(title="Customer_Data_Sample")
    ws4.views.sheetView[0].showGridLines = True
    sample_df = df.head(500)
    for c_idx, col_name in enumerate(sample_df.columns, 1):
        cell = ws4.cell(row=1, column=c_idx, value=col_name)
        cell.fill = navy_fill
        cell.font = header_font

    for r_idx, row in enumerate(sample_df.values, 2):
        for c_idx, val in enumerate(row, 1):
            ws4.cell(row=r_idx, column=c_idx, value=val).border = thin_border

    # Adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"[+] Customer retention Excel workbook created at: {output_path}")


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    clean_csv = os.path.join(base_dir, 'data', 'cleaned_customer_churn_data.csv')
    wb_out = os.path.join(base_dir, 'dashboards', 'customer_retention_workbook.xlsx')
    build_retention_excel_workbook(clean_csv, wb_out)
